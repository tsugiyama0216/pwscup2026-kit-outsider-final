"""Retrospective team24 diagnosis. Reads original inputs; saves aggregates only.

Run from any directory using the bundled Python with numpy/pandas/openpyxl.
No new synthesis, attack submission, private-row export, or privacy-budget use.
"""
from pathlib import Path
import ast
import hashlib
import io
import json
import sys
import zipfile

import numpy as np
import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
SOURCES = REPO.parent / "pwscup2026-results-publication/prelim/results/20260909/sources"
DOWNLOADS = Path.home() / "Downloads"
ACTUAL = next(DOWNLOADS.glob("PWSCup2026_prelim_team24_*"))
TARGETS = DOWNLOADS / "PWSCup2026_prelim_attack"
sys.path.insert(0, str(REPO / "prelim/workspace/starter/src"))
from pwscup2026.scoring.attack_score import score_mia, score_aia


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def clean(v):
    if isinstance(v, dict):
        return {str(k): clean(x) for k, x in v.items()}
    if isinstance(v, (list, tuple, np.ndarray)):
        return [clean(x) for x in v]
    if isinstance(v, (np.bool_, bool)):
        return bool(v)
    if isinstance(v, (np.integer, int)):
        return int(v)
    if isinstance(v, (np.floating, float)):
        return float(v) if np.isfinite(v) else None
    return v


def aia_stats(frame, time_hat):
    m = frame.label.eq("member").to_numpy()
    r = frame.is_rare.astype(bool).to_numpy()
    t = frame.time.to_numpy()
    p = np.asarray(time_hat)
    result = score_aia(None, p[m], frame.onset.to_numpy()[m], t[m],
                       None, p[~m], frame.onset.to_numpy()[~m], t[~m],
                       cfg, member_rare_mask=r[m], control_rare_mask=r[~m])
    ans = dict(result.breakdown, score=result.score)
    for label, mask in [('member', m), ('control', ~m), ('member_rare', m & r), ('control_rare', ~m & r)]:
        success = np.abs(p - t) <= .5
        ans[label + '_n'] = int(mask.sum())
        ans[label + '_correct'] = int(success[mask].sum())
        ans[label + '_horizon_truth_n'] = int((mask & (t >= 9.5)).sum())
        ans[label + '_horizon_correct'] = int((mask & (t >= 9.5) & success).sum())
        ans[label + '_early_correct'] = int((mask & (t < 9.5) & success).sum())
    horizon_guess = np.abs(p - 10.) <= .5
    ans['predicted_horizon_n'] = int(horizon_guess.sum())
    ans['horizon_guesses_for_members'] = int((horizon_guess & m).sum())
    ans['horizon_guesses_for_controls'] = int((horizon_guess & ~m).sum())
    ans['predictions_above_10_5'] = int((p > 10.5).sum())
    ans['prediction_unique_count'] = int(pd.Series(p).nunique())
    ans['prediction_in_C_times'] = int(np.isin(p, C.time.to_numpy()).sum())
    ans['prediction_near_C_times_1e-8'] = int((np.min(np.abs(p[:, None] - C.time.to_numpy()[None, :]), axis=1) < 1e-8).sum())
    return ans


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    # Hash gate: do not silently substitute practice data for the actual world.
    tree = ast.parse((REPO / 'prelim/workspace/attack_phase_v8/actual_inputs.py').read_text(encoding='utf-8'))
    expected = next(ast.literal_eval(n.value) for n in tree.body if isinstance(n, ast.Assign)
                    and any(isinstance(t, ast.Name) and t.id == 'EXPECTED_ACTUAL_HASHES' for t in n.targets))
    for name, h in expected.items():
        assert digest(ACTUAL / name).upper() == h, name
    workbook = SOURCES / '公開用結果表_20260909.xlsx'
    wb = openpyxl.load_workbook(workbook, data_only=True, read_only=True)
    downloaded_wb = openpyxl.load_workbook(DOWNLOADS / workbook.name, data_only=True, read_only=True)
    assert wb.sheetnames == downloaded_wb.sheetnames
    assert all(list(wb[n].values) == list(downloaded_wb[n].values) for n in wb.sheetnames)
    def sheet(name):
        rows = list(wb[name].values)
        return pd.DataFrame(rows[1:], columns=rows[0])
    summary = sheet('総括')
    own = summary.loc[summary.iloc[:, 0].eq(24)].iloc[0]
    meta = dict(list(wb['メタ'].values)[1:])
    mia_official = sheet('攻撃_標的別_MIA')
    mia_official = mia_official[mia_official.iloc[:, 1].eq(24)].set_index('攻撃者')
    aia_official = sheet('攻撃_標的別_AIA')
    aia_official = aia_official[aia_official.iloc[:, 1].eq(24)].set_index('攻撃者')
    cfg = {'scoring': {'low_fpr': meta['config.scoring.low_fpr'], 'aia_rare_se_z': meta['config.scoring.aia_rare_se_z']},
           'aia': {'tau_time': meta['config.aia.tau_time']}}
    truthdir = SOURCES / 'ground_truth'
    pool = pd.read_csv(truthdir / 'pool_ids_24.csv')
    truth = pd.read_csv(truthdir / 'aia_truth_24.csv').sort_values('challenge_row_id')
    B = pd.read_csv(ACTUAL / 'B.csv')
    Bself = pd.read_csv(ACTUAL / 'B_self.csv')
    A = pd.read_csv(ACTUAL / 'A_bg.csv')
    Cpath = REPO / 'prelim/anonymization/anonymization-submitted-v5/C.csv'
    assert digest(Cpath) == meta['sha256.defense.C_24.csv'] == digest(TARGETS / 'C_24.csv')
    C = pd.read_csv(Cpath)
    assert pool.record_id.is_unique and B.record_id.is_unique
    B = B.merge(pool, on='record_id', validate='one_to_one')
    assert len(B) == 1049
    numeric = ['age', 'BMI', 'SBP', 'TG', 'HDL', 'ALT', 'FPG']
    cov = [c for c in C.columns if c not in ['time', 'onset', 'death']]
    scale = A[numeric].std().to_numpy()
    differences = (B[numeric].to_numpy()[:, None, :] - C[numeric].to_numpy()[None, :, :]) / scale
    distances = np.sqrt(np.einsum('ijk,ijk->ij', differences, differences))
    B['nearest_C_numeric7'] = distances.min(axis=1)
    all_detected_rare = set()
    rare_seen = set()
    mia_rows, aia_rows, rare_details = [], [], []
    for attacker in sorted(mia_official.index):
        zpath = SOURCES / f'attack_submissions/cohort{attacker:02d}_F.zip'
        with zipfile.ZipFile(zpath) as z:
            files = {Path(n).name: n for n in z.namelist()}
            mia_bytes, aia_bytes = z.read(files['F_mia.csv']), z.read(files['F_aia.csv'])
            assert hashlib.sha256(mia_bytes).hexdigest() == meta[f'sha256.attack.{attacker}.F_mia.csv']
            assert hashlib.sha256(aia_bytes).hexdigest() == meta[f'sha256.attack.{attacker}.F_aia.csv']
            f = pd.read_csv(io.BytesIO(mia_bytes))
            a = pd.read_csv(io.BytesIO(aia_bytes))
        attacker_pool = pd.read_csv(truthdir / f'pool_ids_{attacker}.csv')
        f = f.merge(attacker_pool, on='record_id', validate='one_to_one')
        assert len(f) == 1049
        y = f.pool_id.isin(pool.pool_id).to_numpy()
        rare = f.is_rare.astype(bool).to_numpy()
        conf = f['24'].to_numpy()
        result = score_mia(conf, y, cfg, rare_mask=rare)
        off = mia_official.loc[attacker]
        assert np.isclose(result.score, off['t_j（悪いほう）'], atol=1e-12)
        assert np.isclose(result.breakdown['tpr_at_low_fpr'], off['TPR（全体）'], atol=1e-12)
        assert np.isclose(result.breakdown['auc'], off['AUC'], atol=1e-12)
        mrow = dict(attacker=int(attacker), **result.breakdown, members=int(y.sum()),
                    rare_members=int((y & rare).sum()), rare_nonmembers=int((~y & rare).sum()),
                    unique_confidences=int(pd.Series(conf).nunique()))
        # Strict zero-false-positive detections: distinct from interpolated TPR@.01.
        if (rare & ~y).any():
            cutoff = conf[rare & ~y].max()
            detected = set(f.loc[rare & y & (conf > cutoff), 'pool_id'])
            all_detected_rare |= detected
            rare_seen |= set(f.loc[rare & y, 'pool_id'])
            mrow['rare_members_above_all_rare_nonmembers'] = len(detected)
            member_ranks = []
            for idx in np.where(rare & y)[0]:
                member_ranks.append({'strictly_higher_rare_nonmembers': int((conf[rare & ~y] > conf[idx]).sum()),
                                     'tied_rare_nonmembers': int((conf[rare & ~y] == conf[idx]).sum())})
            mrow['rare_member_ranks'] = member_ranks
        mia_rows.append(mrow)
        a = a[a.target_k.eq(24)].merge(truth, on='challenge_row_id', validate='one_to_one').sort_values('challenge_row_id')
        assert len(a) == 326
        arow = dict(attacker=int(attacker), **aia_stats(a, a.time_hat.to_numpy()))
        ao = aia_official.loc[attacker]
        for k, column in [('score', '失点（悪いほう）'), ('r_time_all', 'R（全体）'),
                          ('r_time_rare', 'R（希少な患者だけ）'), ('succ_time_member', 'p_m'), ('succ_time_control', 'p_c')]:
            assert np.isclose(arow[k], ao[column], atol=1e-12), (attacker, k)
        aia_rows.append(arow)
    mframe, aframe = pd.DataFrame(mia_rows), pd.DataFrame(aia_rows)
    raw = mframe.score.mean()
    chance = float(own['beta_chance'])
    fpr = cfg['scoring']['low_fpr']
    beta = fpr + max(raw - chance, 0.)
    loss = np.clip((beta - fpr) / ((meta['config.scoring.mia_fold_K'] - 1) * fpr), 0., 1.)
    assert np.isclose(raw, own['beta_raw'], atol=1e-12)
    assert np.isclose(beta, own['beta'], atol=1e-12)
    assert np.isclose(loss, own['MIA 成分'], atol=1e-12)
    assert np.isclose(aframe.score.max(), own['AIA 成分'], atol=1e-12)
    # No private rows or direct record identifiers are written below.
    def group_distance(mask):
        d = B.loc[mask, 'nearest_C_numeric7'].to_numpy()
        return {'n': len(d), 'min': d.min() if len(d) else None,
                'median': np.median(d) if len(d) else None, 'max': d.max() if len(d) else None}
    rare = B.is_rare.astype(bool)
    groups = {'detected_rare': group_distance(B.pool_id.isin(all_detected_rare)),
              'eligible_undetected_rare': group_distance(B.pool_id.isin(rare_seen - all_detected_rare)),
              'all_rare': group_distance(rare), 'nonrare': group_distance(~rare)}
    def rows_match_count(left, right, columns, decimals=None):
        left_values = (left[columns] if decimals is None else left[columns].round(decimals)).to_numpy()
        right_values = (right[columns] if decimals is None else right[columns].round(decimals)).to_numpy()
        keys = set(map(tuple, right_values))
        return sum(tuple(row) in keys for row in left_values)
    copying = {name: {'C_to_B': rows_match_count(C, B, columns), 'C_to_A_bg': rows_match_count(C, A, columns)}
               for name, columns in [('all13', list(C.columns)), ('covariates10', cov),
                                     ('numeric7', numeric), ('outcomes3', ['time', 'onset', 'death'])]}
    copying['each_numeric_value_from_A_bg'] = {col: int(C[col].isin(A[col]).sum()) for col in numeric}
    copying['rounded_to_6_decimals'] = {
        name: {'C_to_B': rows_match_count(C, B, columns, 6), 'C_to_A_bg': rows_match_count(C, A, columns, 6)}
        for name, columns in [('all13', list(C.columns)), ('covariates10', cov),
                              ('numeric7', numeric), ('outcomes3', ['time', 'onset', 'death'])]}
    parentpath = REPO.parent / 'pwscup2026-kit/local-output/prelim/team24/opendp_aim/release-001/C.csv'
    parent_comparison = {'available': parentpath.is_file()}
    if parentpath.is_file():
        parent = pd.read_csv(parentpath)
        parent_comparison.update(covariates_exactly_unchanged=bool(np.array_equal(parent[cov].to_numpy(), C[cov].to_numpy())),
                                 outcome_rows_changed_exact_float=int((parent[['time','onset','death']] != C[['time','onset','death']]).any(axis=1).sum()),
                                 outcome_rows_changed_tolerance_1e_8=int((~np.isclose(parent[['time','onset','death']], C[['time','onset','death']], atol=1e-8, rtol=0)).any(axis=1).sum()),
                                 onset_rows_changed=int(parent.onset.ne(C.onset).sum()),
                                 death_rows_changed=int(parent.death.ne(C.death).sum()))
    horizon = aia_stats(truth, np.full(len(truth), 10.))
    # Fixed diagnostic recipes, evaluated retrospectively. No label-based fitting.
    # A cohort substitution changes only the released C, keeping target24 truth/QI.
    challenge = pd.read_csv(TARGETS / 'aia_challenge_24.csv').merge(truth, on='challenge_row_id', validate='one_to_one').sort_values('challenge_row_id')
    qcols = ['age', 'SBP', 'TG', 'HDL', 'ALT']
    def transform(frame):
        x = frame[qcols].to_numpy().copy()
        x[:, 2] = np.log(x[:, 2])
        x[:, 4] = np.log(x[:, 4])
        return x
    qscale = transform(A).std(axis=0, ddof=1)
    qx = transform(challenge) / qscale
    def nn3(frame):
        ref = transform(frame) / qscale
        d2 = (qx * qx).sum(axis=1)[:, None] + (ref * ref).sum(axis=1)[None, :] - 2 * qx @ ref.T
        return np.maximum(np.partition(d2, 2, axis=1)[:, :3], 0).mean(axis=1)
    def selected_stats(scores):
        order = np.lexsort((challenge.challenge_row_id.to_numpy(), -np.asarray(scores)))
        p = np.full(len(challenge), 11.)
        p[order[:163]] = 10.
        s = aia_stats(challenge, p)
        return {k:s[k] for k in ['score','r_time_all','r_time_rare','member_correct','control_correct','horizon_guesses_for_members','horizon_guesses_for_controls']}
    bgdist = nn3(A)
    aids = sorted(int(p.stem.split('_')[1]) for p in TARGETS.glob('C_*.csv'))
    substitutions = []
    for cid in aids:
        c = pd.read_csv(TARGETS / f'C_{cid}.csv')
        assert digest(TARGETS / f'C_{cid}.csv') == meta[f'sha256.defense.C_{cid}.csv']
        d = nn3(c)
        s = selected_stats(np.log(bgdist + 1e-12) - np.log(d + 1e-12))
        substitutions.append(dict(release_cohort=cid, **s))
    current_sub = next(r for r in substitutions if r['release_cohort']==24)
    own_pref = C.prefecture.value_counts(normalize=True)
    bg_pref = A.prefecture.value_counts(normalize=True)
    prefscore = np.log(challenge.prefecture.map(own_pref).fillna(1e-12).to_numpy()) - np.log(challenge.prefecture.map(bg_pref).fillna(1e-12).to_numpy())
    diagnostic_selection = {
        'method': 'Select top163 by five-numeric-QI 3NN distance ratio to C versus A_bg; predict10, others11. TG/ALT log; A_bg std scaling. Fixed recipe, retrospective, not attacker-code reproduction.',
        'with_C24': current_sub,
        'public_A_bg_density_only': selected_stats(-bgdist),
        'C24_prefecture_frequency_ratio_only': selected_stats(prefscore),
        'C_substitutions': substitutions,
        'C24_rank_among_substitutions': 1 + sum(r['score'] > current_sub['score'] + 1e-12 for r in substitutions),
        'warning': 'No hyperparameter search; disclosed truth used only for evaluation. One challenge set and one heuristic cannot prove or exclude actual leakage.'}
    toy_path = REPO.parent / 'pwscup2026-kit/local-output/opendp_aim/toy-attack-v5-uspec/toy_attack_report.json'
    toy = json.loads(toy_path.read_text(encoding='utf-8'))
    earlier_validation = {
        'report': 'local-output/opendp_aim/toy-attack-v5-uspec/toy_attack_report.json',
        'C_is_actual_submitted_C': toy['inputs']['c_sha256'] == digest(Cpath),
        'B_is_actual_B': toy['inputs']['b_practice_sha256'] == digest(ACTUAL / 'B.csv'),
        'A_bg_is_actual_A_bg': toy['inputs']['a_bg_sha256'] == digest(ACTUAL / 'A_bg.csv'),
        'warning': 'The stored v5 toy report evaluates a practice-world C, not the actual submitted C.'}
    evaluated = summary[summary['Anon'].notna()]
    output = {
        'scope': 'Team24 preliminary submitted anonymization v5; retrospective diagnosis with released truth; no new DP release',
        'source_checks': {'actual_input_hashes_match': True, 'all_workbook_cell_values_match_download': True,
                          'workbook_bytes_match_download': digest(workbook) == digest(DOWNLOADS / workbook.name),
                          'submitted_C_matches_official_hash': True, '20_attack_CSV_hash_pairs_match': True,
                          'all_20_MIA_and_AIA_scores_match': True,
                          'chance_calibration': 'Official beta_chance reused; 20000-permutation calibration not rerun'},
        'summary': own.to_dict(),
        'anonymity_rank': int(evaluated.Anon.gt(own.Anon).sum() + 1),
        'aia_loss_rank_ascending': int(evaluated['AIA 成分'].lt(own['AIA 成分'] - 1e-12).sum() + 1),
        'MIA': {'attackers': mia_rows, 'raw_mean': raw,
                'all_candidates_mean_TPR': mframe.tpr_at_low_fpr.mean(),
                'rare_worst_case_increment': raw - mframe.tpr_at_low_fpr.mean(),
                'all_candidates_zero_attackers': int(mframe.tpr_at_low_fpr.eq(0).sum()),
                'rare_eligible_unique_people': len(rare_seen), 'rare_strict_detected_unique_people': len(all_detected_rare),
                'raw_contribution_by_positive_attacker': {int(r.attacker): float(r.score / 20) for r in mframe.itertuples() if r.score > 0},
                'beta_chance_official': chance, 'beta_recomputed': beta, 'loss_recomputed': loss,
                'rare_distance_diagnostics': groups},
        'AIA': {'attackers': aia_rows, 'constant_10_year_baseline': horizon,
                'selection_diagnostics': diagnostic_selection},
        'copying_diagnostics': copying,
        'parent_to_v5': parent_comparison,
        'earlier_toy_validation': earlier_validation,
        'limitations': ['Published F files contain predictions, not attacker training code.',
                        'One realized DP release; cannot attribute a causal effect to a generation stage.',
                        'The known-truth diagnostic is retrospective, not held-out validation.',
                        'Attackers share the same released C and sampling world; independent-trial uncertainty is not assumed.']}
    (HERE / 'results.json').write_text(json.dumps(clean(output), ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(clean({k: v for k,v in output.items() if k not in ['MIA','AIA']}), ensure_ascii=False, indent=2))
    print('MIA', mframe[['attacker','tpr_at_low_fpr','tpr_rare','score','members','rare_members','rare_members_above_all_rare_nonmembers']].to_string(index=False))
    print('AIA', aframe[['attacker','score','member_correct','control_correct','member_horizon_correct','control_horizon_correct','member_early_correct','control_early_correct','predicted_horizon_n','prediction_near_C_times_1e-8']].to_string(index=False))
    print('MIA_AGGREGATE', json.dumps(clean({k:v for k,v in output['MIA'].items() if k != 'attackers'}), ensure_ascii=False))
    print('HORIZON_BASELINE', json.dumps(clean(horizon), ensure_ascii=False))
    print('SELECTION_DIAGNOSTICS', json.dumps(clean(diagnostic_selection), ensure_ascii=False))
